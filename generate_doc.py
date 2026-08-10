import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_table_rtl(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        bidiVisual = OxmlElement('w:bidiVisual')
        tblPr[0].append(bidiVisual)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def make_run_rtl(run):
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)

def build_schema_docx():
    doc = Document()

    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("الفصل الرابع: تصميم قاعدة البيانات للنظام (NexusPay Store)")
    run.font.name = 'Traditional Arabic'
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42) # Slate-900
    make_run_rtl(run)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("مواصفات جداول قاعدة البيانات، حقولها، أنواع البيانات، المفاتيح والوصف التفصيلي")
    sub_run.font.name = 'Traditional Arabic'
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(71, 85, 105)
    make_run_rtl(sub_run)

    doc.add_paragraph() # Spacer

    tables_data = [
        {
            "num": "1.3.4",
            "title": "جدول المستخدمين:",
            "caption": "جدول (1-4) المستخدمين (users)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "الرقم التعريفي الموحد للمستخدم (مفتاح أساسي)"),
                ("email", "VARCHAR(255)", "", "البريد الإلكتروني الحقيقي للمستخدم (فريد)"),
                ("password_hash", "VARCHAR(255)", "", "كلمة المرور المشفرة بتقنية PBKDF2 (Salt:Hash)"),
                ("full_name", "VARCHAR(100)", "", "الاسم الكامل للمستخدم (الاسم في البطاقة/المنصة)"),
                ("phone", "VARCHAR(30)", "", "رقم هاتف الواتساب / الجوال والتواصل"),
                ("phone_number", "VARCHAR(30)", "", "رقم الهاتف المباشر الإضافي"),
                ("country_code", "VARCHAR(10)", "", "رمز الدولة للمستخدم (افتراضي: +967)"),
                ("role_code", "VARCHAR(50)", "FK", "رمز دور الحساب المرتبط بجدول الأدوار (roles)"),
                ("status", "VARCHAR(30)", "", "حالة الحساب (ACTIVE, SUSPENDED, PENDING_VERIFICATION)"),
                ("is_2fa_enabled", "BOOLEAN", "", "تفعيل المصادقة الثنائية بالحساب (True/False)"),
                ("deleted_at", "TIMESTAMPTZ", "", "تاريخ الحذف اللطيف (Soft Delete)"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت إنشاء الحساب في النظام"),
                ("updated_at", "TIMESTAMPTZ", "", "تاريخ ووقت آخر تحديث لبيانات الحساب")
            ]
        },
        {
            "num": "2.3.4",
            "title": "جدول الأدوار:",
            "caption": "جدول (2-4) الأدوار (roles)",
            "fields": [
                ("code", "VARCHAR(50)", "PK", "رمز الدور الوظيفي (ADMIN, USER) (مفتاح أساسي)"),
                ("title", "VARCHAR(100)", "", "المسمى الوظيفي والدور باللغتين العربية والإنجليزيّة"),
                ("description", "TEXT", "", "وصف مفصل للصلاحيات والمسؤوليات الممنوحة للدور")
            ]
        },
        {
            "num": "3.3.4",
            "title": "جدول الأذونات والصلاحيات:",
            "caption": "جدول (3-4) الأذونات (permissions)",
            "fields": [
                ("code", "VARCHAR(100)", "PK", "رمز الإذن الفريد للنظام (مفتاح أساسي)"),
                ("module", "VARCHAR(50)", "", "اسم وحدة النظام التابع لها الإذن (ORDERS, USERS...)")
            ]
        },
        {
            "num": "4.3.4",
            "title": "جدول أذونات الأدوار:",
            "caption": "جدول (4-4) أذونات الأدوار (role_permissions)",
            "fields": [
                ("role_code", "VARCHAR(50)", "PK, FK", "رمز الدور المرتبط بجدول الأدوار (roles)"),
                ("permission_code", "VARCHAR(100)", "PK, FK", "رمز الإذن المرتبط بجدول الأذونات (permissions)")
            ]
        },
        {
            "num": "5.3.4",
            "title": "جدول جلسات المستخدمين والأجهزة:",
            "caption": "جدول (5-4) جلسات المستخدمين (user_sessions)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف الجلسة الفريد (مفتاح أساسي)"),
                ("user_id", "VARCHAR(35)", "FK", "رقم المستخدم المرتبط بالجلسة بجدول (users)"),
                ("access_token_hash", "VARCHAR(255)", "", "تشفير توكن الوصول (Access Token Hash)"),
                ("refresh_token_hash", "VARCHAR(255)", "", "تشفير توكن التحديث (Refresh Token Hash)"),
                ("device_fingerprint", "VARCHAR(255)", "", "بصمة الجهاز المتصل والمطابقة الأمنيّة"),
                ("ip_address", "VARCHAR(45)", "", "عنوان بروتوكول الإنترنت (IP Address)"),
                ("expires_at", "TIMESTAMPTZ", "", "تاريخ ووقت انتهاء صلاحية الجلسة")
            ]
        },
        {
            "num": "6.3.4",
            "title": "جدول المحافظ المالية للمستخدمين:",
            "caption": "جدول (6-4) المحافظ المالية (wallets)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف المحفظة المالية (مفتاح أساسي)"),
                ("user_id", "VARCHAR(35)", "FK", "رقم المستخدم صاحب المحفظة المرتبط بجدول (users)"),
                ("available_balance", "DECIMAL(16,4)", "", "الرصيد المتاح المباشر للشراء والتسوق"),
                ("held_balance", "DECIMAL(16,4)", "", "الرصيد المحتجز قيد تنفيذ الطلبات"),
                ("currency", "VARCHAR(10)", "", "رمز عملة المحفظة (افتراضي: USD)")
            ]
        },
        {
            "num": "7.3.4",
            "title": "جدول طلبات وحركات إيداع المحفظة:",
            "caption": "جدول (7-4) طلبات إيداع المحفظة (wallet_topup_transactions)",
            "fields": [
                ("id", "VARCHAR(50)", "PK", "رقم عملية الإيداع الفريد (مفتاح أساسي)"),
                ("user_id", "VARCHAR(35)", "FK", "رقم المستخدم المودِع المرتبط بجدول (users)"),
                ("amount_usd", "DECIMAL(16,4)", "", "المبلغ المطلوب إيداعه بالدولار الأمريكي"),
                ("payment_method", "VARCHAR(50)", "", "طريقة التحويل المالي (الكريمي، جيب، ون كاش...)"),
                ("payment_reference", "VARCHAR(255)", "", "رقم المرجع / الإشعار الوارد من العميل"),
                ("status", "VARCHAR(30)", "", "حالة الطلب (PENDING, COMPLETED, FAILED, REFUNDED)"),
                ("gateway_response", "JSONB", "", "البيانات والإثباتات المرفقة مع حركة الإيداع"),
                ("country_code", "VARCHAR(10)", "", "رمز دولة الإيداع المالي (YE)"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت تقديم طلب الإيداع"),
                ("completed_at", "TIMESTAMPTZ", "", "تاريخ ووقت تأكيد إضافة الرصيد للمحفظة")
            ]
        },
        {
            "num": "8.3.4",
            "title": "جدول دفتر قيود المحفظة:",
            "caption": "جدول (8-4) دفتر قيود المحفظة (wallet_ledgers)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف القيد المالي الفريد (مفتاح أساسي)"),
                ("wallet_id", "VARCHAR(35)", "FK", "رقم المحفظة المرتبطة بجدول (wallets)"),
                ("correlation_id", "VARCHAR(100)", "", "رمز تتبع العملية الموحد لضمان التدقيق"),
                ("transaction_direction", "VARCHAR(10)", "", "اتجاه الحركة المالية (DEBIT / CREDIT)"),
                ("entry_type", "VARCHAR(20)", "", "نوع القيد المالي (HOLD, CAPTURE, RELEASE, CREDIT, DEBIT)"),
                ("amount", "DECIMAL(16,4)", "", "مبلغ الحركة المالية بأسلوب القيد المزدوج"),
                ("balance_before", "DECIMAL(16,4)", "", "رصيد المحفظة قبل إجراء العملية"),
                ("balance_after", "DECIMAL(16,4)", "", "رصيد المحفظة بعد إجراء العملية"),
                ("reference_type", "VARCHAR(50)", "", "مصدر المرجع (ORDER, DEPOSIT, REFUND)"),
                ("reference_id", "VARCHAR(100)", "", "رقم المرجع المرتبط بالعملية"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت تسجيل القيد المحاسبي")
            ]
        },
        {
            "num": "9.3.4",
            "title": "جدول الفئات والتصنيفات الرئيسية:",
            "caption": "جدول (9-4) الفئات والتصنيفات (categories)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف الفئة الفريد (مفتاح أساسي)"),
                ("slug", "VARCHAR(100)", "", "المعرف المختصر الفريد للرابط (Slug)"),
                ("title_json", "JSONB", "", "اسم التصنيف باللغات المعتمدة (عربي / إنجليزي)"),
                ("image_url", "TEXT", "", "رابط صورة الشعار / الأيقونة الخاصة بالتصنيف"),
                ("sort_order", "INT", "", "ترتيب ظهور الفئة في الواجهة الرئيسية"),
                ("is_active", "BOOLEAN", "", "حالة تفعيل التصنيف على المنصة (True/False)"),
                ("deleted_at", "TIMESTAMPTZ", "", "تاريخ الحذف اللطيف للتصنيف")
            ]
        },
        {
            "num": "10.3.4",
            "title": "جدول المنتجات والباقات الرقمية:",
            "caption": "جدول (10-4) المنتجات والباقات (products)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف المنتج (مثل: pubg-mobile) (مفتاح أساسي)"),
                ("category_id", "VARCHAR(35)", "FK", "معرف التصنيف التابع له بجدول (categories)"),
                ("title_json", "JSONB", "", "اسم الخدمة/اللعبة باللغات المعتمدة"),
                ("slug", "VARCHAR(150)", "", "رابط اسم المنتج الفريد بالمتصفح"),
                ("product_type", "VARCHAR(50)", "", "نوع الخدمة (DIRECT_TOPUP, VOUCHER_CODE, SUBSCRIPTION...)"),
                ("base_cost", "DECIMAL(16,4)", "", "تكلفة المنتج الأساسية من المورد"),
                ("profit_margin_percent", "DECIMAL(5,2)", "", "نسبة هامش الربح المضافة على المنتج"),
                ("sale_price", "DECIMAL(16,4)", "", "سعر البيع النهائي للعميل بالدولار"),
                ("image_url", "TEXT", "", "رابط غلاف وبوستر المنتج الرقمي"),
                ("country_code", "VARCHAR(10)", "", "نطاق الدولة المدعومة (GLOBAL, YE...)"),
                ("packages_json", "JSONB", "", "قائمة الباقات والفئات المتاحة للشحن"),
                ("currency", "VARCHAR(10)", "", "رمز العملة (USD)"),
                ("is_available", "BOOLEAN", "", "توفر المنتج للبيع والشحن الفوري"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ إضافة المنتج للنظام"),
                ("deleted_at", "TIMESTAMPTZ", "", "تاريخ أرشفة/حذف المنتج")
            ]
        },
        {
            "num": "11.3.4",
            "title": "جدول مخزون الأكواد الرقمية والدفعات:",
            "caption": "جدول (11-4) مخزون الأكواد الرقمية (inventory_stock)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف الكود الرقمي في المخزون (مفتاح أساسي)"),
                ("product_id", "VARCHAR(50)", "FK", "معرف المنتج المرتبط بجدول (products)"),
                ("package_id", "VARCHAR(50)", "", "معرف الباقة الرقمية المخصصة الكود"),
                ("code_payload", "TEXT", "", "كود البطاقة/التفعيل المشفر بطبقة AES-256"),
                ("serial_number", "VARCHAR(100)", "", "الرقم التسلسلي للبطاقة الرقمية (Serial No)"),
                ("status", "VARCHAR(30)", "", "حالة الكود (AVAILABLE, RESERVED, SOLD, EXPIRED)"),
                ("supplier_id", "VARCHAR(50)", "", "رمز المورد/المزود المورّد للكود"),
                ("batch_number", "VARCHAR(100)", "", "رقم دفعة التوريد الاستيرادية"),
                ("cost_price", "DECIMAL(16,4)", "", "سعر شراء التكلفة للكود الفردي"),
                ("purchase_reference", "VARCHAR(100)", "", "مرجع فاتورة الشراء والتوريد"),
                ("order_id", "VARCHAR(35)", "FK", "رقم الطلب المباع له الكود (عند البيع)"),
                ("expires_at", "TIMESTAMPTZ", "", "تاريخ انتهاء صلاحية الكود الرقمي"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ إدخال الكود لمخزون النظام")
            ]
        },
        {
            "num": "12.3.4",
            "title": "جدول طلبات الشراء الرئيسية:",
            "caption": "جدول (12-4) طلبات الشراء الرئيسية (orders)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "رقم الطلب الرئيسي (مثل: NP-739607) (مفتاح أساسي)"),
                ("order_number", "VARCHAR(50)", "", "الرقم المرجعي الفني للطلب (ORD-XXXXX)"),
                ("correlation_id", "VARCHAR(100)", "", "معرف التتبع الموحد لعملية الشراء"),
                ("idempotency_key", "VARCHAR(100)", "", "مفتاح منع تكرار تنفيذ الشراء بالخطأ"),
                ("user_id", "VARCHAR(35)", "FK", "رقم المستخدم صاحب الطلب المرتبط بـ (users)"),
                ("is_guest", "BOOLEAN", "", "هل الطلب تمت كتابته كعميل زائر دون تسجيل؟"),
                ("guest_email", "VARCHAR(255)", "", "البريد الإلكتروني للعميل الزائر استلام الكود"),
                ("guest_phone", "VARCHAR(45)", "", "رقم جوال العميل الزائر للتواصل والتأكيد"),
                ("user_name", "VARCHAR(255)", "", "اسم العميل المسجل/الزائر صاحب الطلب"),
                ("user_email", "VARCHAR(255)", "", "البريد الإلكتروني المعتمد لإرسال الكود"),
                ("product_id", "VARCHAR(35)", "", "معرف المنتج الرئيسي المطلوب شحنه"),
                ("product_title", "VARCHAR(255)", "", "عنوان المنتج/اللعبة (مثل: شدات ببجي موبايل)"),
                ("package_title", "VARCHAR(255)", "", "اسم الفئة/الباقة المحددة للشحن"),
                ("target_account_id", "VARCHAR(255)", "", "معرف الحساب / ID اللاعب المطلوب الشحن له"),
                ("status", "VARCHAR(40)", "", "حالة الطلب (PENDING, PAYMENT_CONFIRMED, COMPLETED, REJECTED)"),
                ("total_amount", "DECIMAL(16,4)", "", "إجمالي سعر الطلب المدفوع بالدولار"),
                ("payment_method", "VARCHAR(50)", "", "طريقة السداد (WALLET, YEMEN_TRANSFER)"),
                ("delivered_payload", "JSONB", "", "بيانات الكود والإيصال المسلم للعميل"),
                ("items_json", "JSONB", "", "تفاصيل جميع الباقات المطلوبة في حال السلة المتعددة"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت تقديم طلب الشراء"),
                ("updated_at", "TIMESTAMPTZ", "", "تاريخ ووقت آخر تحديث لحالة الطلب")
            ]
        },
        {
            "num": "13.3.4",
            "title": "جدول تفاصيل عناصر الطلب:",
            "caption": "جدول (13-4) تفاصيل عناصر الطلب (order_items)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف عنصر الطلب (مفتاح أساسي)"),
                ("order_id", "VARCHAR(35)", "FK", "رقم الطلب الرئيسي المرتبط بجدول (orders)"),
                ("product_id", "VARCHAR(35)", "FK", "رقم المنتج المرتبط بجدول (products)"),
                ("unit_price", "DECIMAL(16,4)", "", "سعر الوحدة الواحدة من الباقة"),
                ("quantity", "INT", "", "الكمية المطلوبة من المنتج/الباقة")
            ]
        },
        {
            "num": "14.3.4",
            "title": "جدول مدفوعات البوابات:",
            "caption": "جدول (14-4) المدفوعات وبوابات الدفع (payments)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف عملية الدفع (مفتاح أساسي)"),
                ("order_id", "VARCHAR(35)", "FK", "رقم الطلب المرتبط بالدفع بجدول (orders)"),
                ("gateway_code", "VARCHAR(50)", "", "رمز بوابة الدفع (WALLET, KURAIMI, PAYPAL)"),
                ("gateway_tx_id", "VARCHAR(100)", "", "رقم العملية الفريد لدى بوابة الدفع"),
                ("amount", "DECIMAL(16,4)", "", "المبلغ المدفوع بالعملة المحددة"),
                ("currency", "VARCHAR(10)", "", "رمز العملة المستخدمة في العملية"),
                ("status", "VARCHAR(30)", "", "حالة عملية الدفع (PENDING, PAID, FAILED, REFUNDED)"),
                ("gateway_response_payload", "JSONB", "", "الاستجابة الفنية الكاملة الواردة من البوابة")
            ]
        },
        {
            "num": "15.3.4",
            "title": "جدول مزودي الخدمات الخارجيين:",
            "caption": "جدول (15-4) مزودي الخدمات الخارجيين (providers)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف المزود الفريد (مفتاح أساسي)"),
                ("name", "VARCHAR(100)", "", "اسم الشركة / المنصة المزودة للخدمة"),
                ("code", "VARCHAR(50)", "", "رمز المزود البرمجي (مثال: RELOADLY)"),
                ("driver_class", "VARCHAR(100)", "", "فئة السائق البرمجي المشغل للمزود"),
                ("base_url", "TEXT", "", "الرابط الأساسي لـ API المزود"),
                ("logo_url", "TEXT", "", "رابط شعار المزود الخارجي"),
                ("balance", "DECIMAL(16,4)", "", "رصيد حساب المنصة المتاح لدى المزود"),
                ("priority_weight", "INT", "", "أولوية التوجيه الذكي للشراء من المزود"),
                ("is_active", "BOOLEAN", "", "تفعيل المزود واستقبال الطلبات منه"),
                ("in_maintenance", "BOOLEAN", "", "وضع الصيانة المؤقت للمزود")
            ]
        },
        {
            "num": "16.3.4",
            "title": "جدول منتجات المزودين الخارجيين:",
            "caption": "جدول (16-4) منتجات المزودين (provider_products)",
            "fields": [
                ("provider_id", "VARCHAR(35)", "PK, FK", "معرف المزود المرتبط بجدول (providers)"),
                ("product_id", "VARCHAR(35)", "PK, FK", "معرف المنتج المرتبط بجدول (products)"),
                ("provider_sku", "VARCHAR(100)", "", "رمز المنتج (SKU) الفريد لدى نظام المزود"),
                ("provider_cost", "DECIMAL(16,4)", "", "تكلفة شراء المنتج من المزود الخارجي"),
                ("provider_stock", "INT", "", "المخزون الفعلي المتاح لدى المزود"),
                ("last_sync_at", "TIMESTAMPTZ", "", "تاريخ آخر مزامنة تلقائية لبيانات المنتج")
            ]
        },
        {
            "num": "17.3.4",
            "title": "جدول قواطع الحماية التلقائية للمزودين:",
            "caption": "جدول (17-4) قواطع الحماية التلقائية (circuit_breaker_states)",
            "fields": [
                ("provider_id", "VARCHAR(35)", "PK, FK", "معرف المزود المرتبط بجدول (providers)"),
                ("state", "VARCHAR(20)", "", "حالة القاطع (CLOSED, OPEN, HALF_OPEN)"),
                ("consecutive_failures", "INT", "", "عدد محاولات الفشل المتتالية للاتصال"),
                ("cooldown_until", "TIMESTAMPTZ", "", "توقيت إعادة محاولة الاتصال بعد فترة التهدئة")
            ]
        },
        {
            "num": "18.3.4",
            "title": "جدول تقييمات وآراء العملاء:",
            "caption": "جدول (18-4) تقييمات وآراء العملاء (user_reviews)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف التقييم الفريد (مفتاح أساسي)"),
                ("user_id", "VARCHAR(35)", "FK", "رقم المستخدم صاحب التقييم بجدول (users)"),
                ("order_id", "VARCHAR(35)", "FK", "رقم الطلب الميمّم المرتبط بالتقييم (orders)"),
                ("product_id", "VARCHAR(35)", "", "معرف المنتج الذي تم تقييمه"),
                ("rating", "INT", "", "التقييم بالنجوم (من 1 إلى 5 نجوم)"),
                ("title", "VARCHAR(200)", "", "عنوان التقييم المختصر (مثل: خدمة ممتازة)"),
                ("comment", "TEXT", "", "نص التعليق ورأي العميل بالتفصيل"),
                ("is_verified_purchase", "BOOLEAN", "", "هل تم الشراء الفعلي من قِبل صاحب التقييم؟"),
                ("is_approved", "BOOLEAN", "", "موافقة الإدارة على إظهار التقييم بالواجهة"),
                ("reviewer_name", "VARCHAR(100)", "", "اسم العميل الظاهر مع التقييم"),
                ("reviewer_country", "VARCHAR(10)", "", "دولة العميل (افتراضي: YE)"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت إضافة التقييم")
            ]
        },
        {
            "num": "19.3.4",
            "title": "جدول الشكاوى ورسائل الدعم الفني:",
            "caption": "جدول (19-4) الشكاوى ورسائل الدعم (user_complaints)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف الرسالة/الشكوى الفريد (مفتاح أساسي)"),
                ("name", "VARCHAR(100)", "", "اسم المراسِل الكامل"),
                ("email", "VARCHAR(255)", "", "البريد الإلكتروني للتواصل ومتابعة الحل"),
                ("subject", "VARCHAR(50)", "", "موضوع الرسالة (TOPUP, ORDER, BUSINESS, GENERAL)"),
                ("message", "TEXT", "", "محتوى ونص الاستفسار أو الشكوى المقدمة"),
                ("status", "VARCHAR(30)", "", "حالة الشكوى (PENDING, IN_PROGRESS, RESOLVED, REJECTED)"),
                ("admin_notes", "TEXT", "", "رد الإدارة وملاحظات الإجراء المتخذ"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت استلام الشكوى")
            ]
        },
        {
            "num": "20.3.4",
            "title": "جدول طلبات الاسترداد المالي:",
            "caption": "جدول (20-4) طلبات الاسترداد المالي (refunds)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف طلب الاسترداد الفريد (مفتاح أساسي)"),
                ("order_id", "VARCHAR(35)", "FK", "رقم الطلب المراد إرجاعه بجدول (orders)"),
                ("user_id", "VARCHAR(35)", "FK", "رقم المستخدم صاحب الطلب بجدول (users)"),
                ("amount", "DECIMAL(16,4)", "", "مبلغ الاسترداد المطلوب بالدولار"),
                ("reason", "TEXT", "", "سبب طلب إلغاء/استرداد قيمة الخدمة"),
                ("requested_by", "VARCHAR(35)", "", "مُصدر طلب الاسترداد (العميل أو النظام)"),
                ("approved_by", "VARCHAR(35)", "", "المسؤول المصادق على إجراء الاسترداد"),
                ("status", "VARCHAR(30)", "", "حالة الاسترداد (PENDING, APPROVED, REJECTED, COMPLETED)"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ طلب وتأكيد الاسترداد")
            ]
        },
        {
            "num": "21.3.4",
            "title": "جدول أسعار صرف العملات:",
            "caption": "جدول (21-4) أسعار صرف العملات (currency_exchange_rates)",
            "fields": [
                ("code", "VARCHAR(10)", "PK", "رمز العملة الدولي (YER, SAR, USD) (مفتاح أساسي)"),
                ("name", "VARCHAR(50)", "", "اسم العملة الرسمي بالعربية/الإنجليزية"),
                ("symbol", "VARCHAR(10)", "", "رمز/علامة العملة (مثل: ر.ي, ر.س, $)"),
                ("rate_to_usd", "DECIMAL(12,6)", "", "سعر التحويل مقابل الدولار الأمريكي"),
                ("last_synced_at", "TIMESTAMPTZ", "", "تاريخ آخر تحديث ومزامنة لسعر الصرف")
            ]
        },
        {
            "num": "22.3.4",
            "title": "جدول الشجرة الحسابية العامة:",
            "caption": "جدول (22-4) الشجرة الحسابية العامة (chart_of_accounts)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف الحساب المحاسبي الفريد (مفتاح أساسي)"),
                ("account_code", "VARCHAR(50)", "", "رقم الحساب المحاسبي الشجري (مثال: 1010)"),
                ("account_name", "VARCHAR(100)", "", "اسم الحساب في النظام المالي الإداري"),
                ("account_type", "VARCHAR(30)", "", "نوع الحساب (ASSET, LIABILITY, EQUITY, REVENUE, EXPENSE)"),
                ("currency", "VARCHAR(10)", "", "عملة الحساب المحاسبي (USD)")
            ]
        },
        {
            "num": "23.3.4",
            "title": "جدول قيود اليومية العامة:",
            "caption": "جدول (23-4) قيود اليومية العامة (journal_entries)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف قيد اليومية (مفتاح أساسي)"),
                ("reference", "VARCHAR(100)", "", "رقم المرجع السندي لقيد اليومية"),
                ("description", "TEXT", "", "البيان والشرح المحاسبي لقيد اليومية"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت تسجيل القيد المحاسبي")
            ]
        },
        {
            "num": "24.3.4",
            "title": "جدول تفاصيل أسطر قيود اليومية:",
            "caption": "جدول (24-4) تفاصيل أسطر قيود اليومية (journal_entry_lines)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف سطر القيد الفريد (مفتاح أساسي)"),
                ("journal_entry_id", "VARCHAR(35)", "FK", "رقم القيد الرئيسي المرتبط بـ (journal_entries)"),
                ("account_id", "VARCHAR(35)", "FK", "رقم الحساب المرتبط بالشجرة (chart_of_accounts)"),
                ("debit_amount", "DECIMAL(16,4)", "", "مبلغ المدين (Debit Amount)"),
                ("credit_amount", "DECIMAL(16,4)", "", "مبلغ الدائن (Credit Amount)"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ تسجيل سطر القيد")
            ]
        },
        {
            "num": "25.3.4",
            "title": "جدول دفتر الحركات المالية الموحد:",
            "caption": "جدول (25-4) دفتر الحركات المالية الموحد (transactions)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف الحركة المالية الموحدة (مفتاح أساسي)"),
                ("type", "VARCHAR(40)", "", "نوع الحركة (PAYMENT, REFUND, WALLET_RECHARGE, COMMISSION...)"),
                ("user_id", "VARCHAR(35)", "FK", "رقم العميل المعني بالحركة بجدول (users)"),
                ("provider_id", "VARCHAR(50)", "FK", "رقم المزود المعني بالحركة بجدول (providers)"),
                ("order_id", "VARCHAR(35)", "FK", "رقم الطلب المربوط بالحركة (orders)"),
                ("amount", "DECIMAL(16,4)", "", "قيمة المبلغ المالي المنفذ بالحركة"),
                ("balance_before", "DECIMAL(16,4)", "", "رصيد الحساب التجاري قبل الحركة"),
                ("balance_after", "DECIMAL(16,4)", "", "رصيد الحساب التجاري بعد الحركة"),
                ("currency", "VARCHAR(10)", "", "عملة المعاملة المالية (USD)"),
                ("exchange_rate_at_order", "DECIMAL(12,6)", "", "سعر التحويل المعتمد وقت إجراء الحركة"),
                ("reference_id", "VARCHAR(100)", "", "الرقم السندي المرجعي للعملية"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت تنفيذ الحركة المالية")
            ]
        },
        {
            "num": "26.3.4",
            "title": "جدول سجلات المراجعة والتدقيق:",
            "caption": "جدول (26-4) سجلات المراجعة والتدقيق (audit_logs)",
            "fields": [
                ("id", "VARCHAR(35)", "PK", "معرف سجل التدقيق الفريد (مفتاح أساسي)"),
                ("actor_id", "VARCHAR(35)", "", "رقم المستخدم/المسؤول الذي قام بالعملية"),
                ("action", "VARCHAR(100)", "", "نوع الإجراء المنفذ (CREATE, UPDATE, DELETE)"),
                ("entity_type", "VARCHAR(100)", "", "اسم الجدول/الكائن الذي تم تعديله"),
                ("entity_id", "VARCHAR(35)", "", "رقم المعرف الفريد للكائن المعدل"),
                ("old_value", "JSONB", "", "البيانات والقيم القديمة قبل التعديل"),
                ("new_value", "JSONB", "", "البيانات والقيم الجديدة بعد التعديل"),
                ("ip_address", "VARCHAR(45)", "", "عنوان IP لجهاز منفذ العملية"),
                ("created_at", "TIMESTAMPTZ", "", "تاريخ ووقت تسجيل حركة المراجعة")
            ]
        }
    ]

    for t_idx, tdata in enumerate(tables_data):
        # Section Heading: e.g. "1.3.4 : جدول المستخدمين:"
        h_para = doc.add_paragraph()
        h_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = h_para.add_run(f"{tdata['num']} : {tdata['title']}")
        h_run.font.name = 'Traditional Arabic'
        h_run.font.size = Pt(14)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(15, 23, 42)
        make_run_rtl(h_run)

        # Table Caption: e.g. "جدول (1-4) المستخدمين"
        c_para = doc.add_paragraph()
        c_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = c_para.add_run(tdata['caption'])
        c_run.font.name = 'Traditional Arabic'
        c_run.font.size = Pt(12)
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(30, 41, 59)
        make_run_rtl(c_run)

        # Create Table with 4 columns:
        # Col 0: اسم الحقل (Rightmost)
        # Col 1: نوع البيانات
        # Col 2: المفتاح
        # Col 3: الوصف (Leftmost)
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_rtl(table)
        set_table_borders(table, color="B0B0B0", sz="4", val="single")

        # Header Row Formatting
        hdr_cells = table.rows[0].cells
        hdr_titles = ["اسم الحقل", "نوع البيانات", "المفتاح", "الوصف"]
        col_widths = [Inches(1.8), Inches(1.5), Inches(0.8), Inches(2.7)]

        for i in range(4):
            hdr_cells[i].width = col_widths[i]
            set_cell_background(hdr_cells[i], "64748B") # Slate-500 Header background
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(hdr_titles[i])
            r.font.name = 'Traditional Arabic'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            make_run_rtl(r)

        # Data Rows Formatting
        for f_name, f_type, f_key, f_desc in tdata['fields']:
            row_cells = table.add_row().cells

            # Apply cell widths
            for i in range(4):
                row_cells[i].width = col_widths[i]

            # Field Name (Col 0) - Center aligned
            p0 = row_cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f_name)
            r0.font.name = 'Calibri'
            r0.font.size = Pt(11)
            r0.font.bold = True
            r0.font.color.rgb = RGBColor(15, 23, 42)

            # Data Type (Col 1) - Center aligned
            p1 = row_cells[1].paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = p1.add_run(f_type)
            r1.font.name = 'Calibri'
            r1.font.size = Pt(10.5)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(51, 65, 85)

            # Key (Col 2) - Center aligned
            p2 = row_cells[2].paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(f_key)
            r2.font.name = 'Calibri'
            r2.font.size = Pt(10.5)
            r2.font.bold = True
            if "PK" in f_key:
                r2.font.color.rgb = RGBColor(225, 29, 72) # Rose-600 for PK
            elif "FK" in f_key:
                r2.font.color.rgb = RGBColor(37, 99, 235) # Blue-600 for FK

            # Description (Col 3) - Right aligned (Arabic text)
            p3 = row_cells[3].paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r3 = p3.add_run(f_desc)
            r3.font.name = 'Traditional Arabic'
            r3.font.size = Pt(11)
            r3.font.color.rgb = RGBColor(30, 41, 59)
            make_run_rtl(r3)

        doc.add_paragraph() # Spacer between tables

    output_path = r"d:\CoinWestie\NexusPay_Database_Tables.docx"
    doc.save(output_path)
    print(f"Word document generated successfully at: {output_path}")

if __name__ == "__main__":
    build_schema_docx()
